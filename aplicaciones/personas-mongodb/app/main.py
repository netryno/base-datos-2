"""
===========================================================================
Python + MongoDB con FastAPI
===========================================================================

Monolito educativo: conexión, colecciones e índices en database_bootstrap;
aquí el CRUD con PyMongo sobre documentos (equivalente al modelo relacional).
"""

from io import BytesIO
from datetime import date, datetime

from pymongo.errors import DuplicateKeyError
from fastapi import FastAPI, HTTPException, Depends, File, UploadFile
from pydantic import BaseModel, Field, ValidationError, field_validator
from typing import Any, Optional
import re
from openpyxl import load_workbook

from app.database_bootstrap import bootstrap_database, get_database, next_sequence

# ============================================================================
# CONEXIÓN (get_db)
# ============================================================================


def get_db():
    """Generador de acceso a la BD. MongoClient se reutiliza (pool interno)."""
    yield get_database()


def _strip_object_id(doc: dict | None) -> dict | None:
    if doc is None:
        return None
    out = dict(doc)
    out.pop("_id", None)
    return out


def _fecha_to_str(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _prepare_doc(doc: dict | None) -> dict | None:
    cleaned = _strip_object_id(doc)
    if cleaned is None:
        return None
    if "fecha_llegada" in cleaned:
        cleaned["fecha_llegada"] = _fecha_to_str(cleaned["fecha_llegada"])
    return cleaned


# ============================================================================
# MODELOS PYDANTIC
# ============================================================================


class PersonaCreate(BaseModel):
    nombre: str = Field(..., min_length=1, max_length=50)
    primer_apellido: str = Field(..., min_length=1, max_length=50)
    segundo_apellido: Optional[str] = Field(None, max_length=50)
    ci: str = Field(..., min_length=1, max_length=20)

    @field_validator("ci")
    @classmethod
    def ci_solo_numeros(cls, v):
        if not re.match(r"^\d+$", v):
            raise ValueError("CI debe contener solo números")
        return v

    @field_validator("nombre", "primer_apellido", "segundo_apellido")
    @classmethod
    def solo_letras(cls, v):
        if v is not None and not re.match(r"^[a-zA-ZáéíóúÁÉÍÓÚñÑ\s]+$", v):
            raise ValueError("Solo se permiten letras y espacios")
        return v.strip() if v else v


class ViajeCreate(BaseModel):
    persona_id: int = Field(..., gt=0, description="ID de persona que viaja")
    pais_id: int = Field(..., gt=0, description="ID del país destino")
    fecha_llegada: date = Field(..., description="Fecha de llegada")

    @field_validator("fecha_llegada")
    @classmethod
    def no_futuro(cls, v):
        if v > date.today():
            raise ValueError("La fecha no puede ser futura")
        return v


# ============================================================================
# APLICACIÓN FASTAPI
# ============================================================================

app = FastAPI(
    title="Python + MongoDB",
    description="Ejemplo mínimo para comprender la conexión a BD no relacional",
    version="1.0",
)


@app.on_event("startup")
def on_startup():
    bootstrap_database()


@app.get("/")
def root():
    return {
        "mensaje": "Tutorial Python + MongoDB (contenedor mongo en bd-mongo)",
        "endpoints": {
            "POST /by-paul/personas": "Crear persona",
            "POST /by-paul/personas/import-excel": "Importar personas desde .xlsx",
            "GET /by-paul/personas": "Listar personas",
            "DELETE /by-paul/personas/{id}": "Eliminar persona + sus viajes",
            "POST /by-paul/viajes": "Crear viaje (necesita persona_id y pais_id)",
            "GET /by-paul/viajes": "Listar viajes",
        },
        "excel_personas": {
            "formato": "Primera fila: encabezados",
            "columnas": ["nombre", "primer_apellido", "segundo_apellido", "ci"],
            "nota": "segundo_apellido es opcional; ci solo dígitos",
        },
    }


# ============================================================================
# CRUD PERSONAS
# ============================================================================


@app.post("/by-paul/personas", status_code=201)
def crear_persona(persona: PersonaCreate, db=Depends(get_db)):
    doc = {
        "persona_id": next_sequence(db, "persona_id"),
        "nombre": persona.nombre,
        "primer_apellido": persona.primer_apellido,
        "segundo_apellido": persona.segundo_apellido,
        "ci": persona.ci,
    }
    try:
        db.personas.insert_one(doc)
    except DuplicateKeyError:
        raise HTTPException(status_code=400, detail=f"CI {persona.ci} ya existe")
    return _prepare_doc(doc)


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().lower().replace(" ", "_")


def _cell_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _parse_personas_workbook(content: bytes) -> tuple[list[PersonaCreate], list[dict[str, Any]]]:
    workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows, None)
    finally:
        workbook.close()

    if not header_row:
        raise HTTPException(status_code=400, detail="El Excel está vacío")

    headers = [_normalize_header(cell) for cell in header_row]
    required = {"nombre", "primer_apellido", "ci"}
    missing = required - set(headers)
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Faltan columnas obligatorias: {', '.join(sorted(missing))}",
        )

    idx = {name: headers.index(name) for name in headers if name}

    personas: list[PersonaCreate] = []
    errores: list[dict[str, Any]] = []
    excel_row = 1

    workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    try:
        for excel_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
                continue

            def col(name: str) -> Optional[str]:
                if name not in idx:
                    return None
                position = idx[name]
                if position >= len(row):
                    return None
                return _cell_text(row[position])

            payload = {
                "nombre": col("nombre"),
                "primer_apellido": col("primer_apellido"),
                "segundo_apellido": col("segundo_apellido"),
                "ci": col("ci"),
            }
            if not payload["nombre"] and not payload["primer_apellido"] and not payload["ci"]:
                continue

            try:
                personas.append(PersonaCreate(**payload))
            except ValidationError as exc:
                errores.append({"fila": excel_row, "detalle": exc.errors()})
    finally:
        workbook.close()

    return personas, errores


@app.post("/by-paul/personas/import-excel")
async def importar_personas_excel(
    archivo: UploadFile = File(..., description="Excel .xlsx con columnas de personas"),
    db=Depends(get_db),
):
    filename = (archivo.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Solo se aceptan archivos .xlsx")

    content = await archivo.read()
    if not content:
        raise HTTPException(status_code=400, detail="El archivo está vacío")

    personas, errores_validacion = _parse_personas_workbook(content)
    if not personas and errores_validacion:
        raise HTTPException(
            status_code=400,
            detail={"mensaje": "Ninguna fila válida", "errores": errores_validacion},
        )

    insertadas = 0
    omitidas = 0
    for persona in personas:
        existing = db.personas.find_one({"ci": persona.ci}, {"_id": 1})
        if existing:
            omitidas += 1
            continue
        doc = {
            "persona_id": next_sequence(db, "persona_id"),
            "nombre": persona.nombre,
            "primer_apellido": persona.primer_apellido,
            "segundo_apellido": persona.segundo_apellido,
            "ci": persona.ci,
        }
        try:
            db.personas.insert_one(doc)
            insertadas += 1
        except DuplicateKeyError:
            omitidas += 1

    return {
        "archivo": archivo.filename,
        "filas_validas": len(personas),
        "insertadas": insertadas,
        "omitidas_duplicado": omitidas,
        "errores_validacion": errores_validacion,
    }


@app.get("/by-paul/personas")
def listar_personas(db=Depends(get_db)):
    cursor = db.personas.find({}, {"_id": 0}).sort("persona_id", 1)
    return list(cursor)


@app.delete("/by-paul/personas/{persona_id}")
def eliminar_persona(persona_id: int, db=Depends(get_db)):
    persona = db.personas.find_one({"persona_id": persona_id}, {"_id": 0})
    if not persona:
        raise HTTPException(status_code=404, detail="Persona no encontrada")

    db.viajes.delete_many({"persona_id": persona_id})
    db.personas.delete_one({"persona_id": persona_id})

    return {
        "mensaje": "Persona eliminada",
        "persona": persona,
        "nota": "Sus viajes también fueron eliminados (equivalente a ON DELETE CASCADE)",
    }


# ============================================================================
# CRUD VIAJES
# ============================================================================


@app.post("/by-paul/viajes", status_code=201)
def crear_viaje(viaje: ViajeCreate, db=Depends(get_db)):
    if not db.personas.find_one({"persona_id": viaje.persona_id}, {"_id": 1}):
        raise HTTPException(status_code=404, detail="Persona no existe")

    if not db.paises.find_one({"pais_id": viaje.pais_id}, {"_id": 1}):
        raise HTTPException(status_code=404, detail="País no existe")

    doc = {
        "viaje_id": next_sequence(db, "viaje_id"),
        "persona_id": viaje.persona_id,
        "pais_id": viaje.pais_id,
        "fecha_llegada": datetime.combine(viaje.fecha_llegada, datetime.min.time()),
    }
    db.viajes.insert_one(doc)
    return _prepare_doc(doc)


@app.get("/by-paul/viajes")
def listar_viajes(db=Depends(get_db)):
    pipeline = [
        {"$sort": {"viaje_id": 1}},
        {
            "$lookup": {
                "from": "personas",
                "localField": "persona_id",
                "foreignField": "persona_id",
                "as": "persona_doc",
            }
        },
        {
            "$lookup": {
                "from": "paises",
                "localField": "pais_id",
                "foreignField": "pais_id",
                "as": "pais_doc",
            }
        },
        {
            "$project": {
                "_id": 0,
                "viaje_id": 1,
                "fecha_llegada": 1,
                "persona": {
                    "$trim": {
                        "input": {
                            "$concat": [
                                {"$ifNull": [{"$arrayElemAt": ["$persona_doc.nombre", 0]}, ""]},
                                " ",
                                {
                                    "$ifNull": [
                                        {"$arrayElemAt": ["$persona_doc.primer_apellido", 0]},
                                        "",
                                    ]
                                },
                            ]
                        }
                    }
                },
                "pais": {"$arrayElemAt": ["$pais_doc.pais_nombre", 0]},
            }
        },
    ]
    rows = list(db.viajes.aggregate(pipeline))
    for row in rows:
        row["fecha_llegada"] = _fecha_to_str(row.get("fecha_llegada"))
    return rows
