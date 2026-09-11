"""
===========================================================================
Python + PostgreSQL con FastAPI
===========================================================================

Este archivo es INTENCIONALMENTE monolítico para que veas todo el flujo
de conexión a base de datos en un solo lugar.

FLUJO DE CONEXIÓN:
1. psycopg2 → Driver estándar para hablar con PostgreSQL.
2. Variables de entorno → HOST, PORT, USER, PASSWORD, DATABASE (.env).
3. get_db() → Función que crea y cierra conexiones.
4. SQL directo → Queries crudos para ver exactamente qué se ejecuta.

DIFERENCIAS CLAVE CON MYSQL:
- Placeholders: %s (igual que psycopg2/mysql-connector).
- SERIAL / IDENTITY en vez de AUTO_INCREMENT.
- RETURNING * en INSERT (PostgreSQL lo soporta nativamente).
- Las foreign keys se respetan por defecto (como InnoDB).
- RealDictCursor para filas como dict (equivalente a dictionary=True).
"""

from io import BytesIO

from psycopg2 import IntegrityError
from psycopg2.extras import RealDictCursor
from fastapi import FastAPI, HTTPException, Depends, File, UploadFile
from pydantic import BaseModel, Field, ValidationError, field_validator
from typing import Any, Optional
from datetime import date
import re
from openpyxl import load_workbook

from app.database_bootstrap import bootstrap_database, get_connection

# ============================================================================
# CONEXIÓN (get_db)
# ============================================================================


def get_db():
    """Generador de conexiones. Se usa con Depends() en FastAPI."""
    conn = get_connection()
    try:
        yield conn
    finally:
        conn.close()


# ============================================================================
# MODELOS PYDANTIC
# ============================================================================

class PersonaCreate(BaseModel):
    nombre: str = Field(..., min_length=1, max_length=50)
    primer_apellido: str = Field(..., min_length=1, max_length=50)
    segundo_apellido: Optional[str] = Field(None, max_length=50)
    ci: str = Field(..., min_length=1, max_length=20)

    @field_validator('ci')
    @classmethod
    def ci_solo_numeros(cls, v):
        if not re.match(r'^\d+$', v):
            raise ValueError('CI debe contener solo números')
        return v

    @field_validator('nombre', 'primer_apellido', 'segundo_apellido')
    @classmethod
    def solo_letras(cls, v):
        if v is not None and not re.match(r'^[a-zA-ZáéíóúÁÉÍÓÚñÑ\s]+$', v):
            raise ValueError('Solo se permiten letras y espacios')
        return v.strip() if v else v


class ViajeCreate(BaseModel):
    persona_id: int = Field(..., gt=0, description="ID de persona que viaja")
    pais_id: int = Field(..., gt=0, description="ID del país destino")
    fecha_llegada: date = Field(..., description="Fecha de llegada")

    @field_validator('fecha_llegada')
    @classmethod
    def no_futuro(cls, v):
        if v > date.today():
            raise ValueError('La fecha no puede ser futura')
        return v


# ============================================================================
# APLICACIÓN FASTAPI
# ============================================================================

app = FastAPI(
    title="Python + PostgreSQL",
    description="Ejemplo mínimo para comprender la conexión a BD",
    version="1.0",
)


@app.on_event("startup")
def on_startup():
    bootstrap_database()


@app.get("/")
def root():
    return {
        "mensaje": "Tutorial Python + PostgreSQL (BD externa: contenedor my-database)",
        "endpoints": {
            "POST /by-paul/personas": "Crear persona",
            "POST /by-paul/personas/import-excel": "Importar personas desde .xlsx",
            "GET /by-paul/personas": "Listar personas",
            "DELETE /by-paul/personas/{id}": "Eliminar persona + sus viajes",
            "POST /by-paul/viajes": "Crear viaje (necesita persona_id y pais_id)",
            "GET /by-paul/viajes": "Listar viajes",
        },
        "excel_personas": {
            "formato": "Primera fila: encabezados",
            "columnas": ["nombre", "primer_apellido", "segundo_apellido", "ci"],
            "nota": "segundo_apellido es opcional; ci solo dígitos",
        },
    }


# ============================================================================
# CRUD PERSONAS
# ============================================================================

@app.post("/by-paul/personas", status_code=201)
def crear_persona(persona: PersonaCreate, db=Depends(get_db)):
    """
    CREAR persona.

    PostgreSQL permite RETURNING * en el INSERT: obtenemos la fila insertada
    en una sola sentencia (sin lastrowid + SELECT).
    """
    cursor = db.cursor(cursor_factory=RealDictCursor)
    try:
        cursor.execute(
            """
            INSERT INTO personas (nombre, primer_apellido, segundo_apellido, ci)
            VALUES (%s, %s, %s, %s)
            RETURNING *
            """,
            (persona.nombre, persona.primer_apellido, persona.segundo_apellido, persona.ci),
        )
        db.commit()
        return cursor.fetchone()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"CI {persona.ci} ya existe")
    finally:
        cursor.close()


@app.post("/by-paul/personas_ps", status_code=201)
def crear_persona_ps(persona: PersonaCreate, db=Depends(get_db)):
    """
    CREAR persona (variante "procedimiento almacenado").

    PostgreSQL soporta CALL a funciones/procedimientos. Aquí mantenemos INSERT
    + SELECT por ci para el mismo comportamiento del tutorial.
    """
    cursor = db.cursor(cursor_factory=RealDictCursor)
    try:
        cursor.execute(
            """
            INSERT INTO personas (nombre, primer_apellido, segundo_apellido, ci)
            VALUES (%s, %s, %s, %s)
            """,
            (persona.nombre, persona.primer_apellido, persona.segundo_apellido, persona.ci),
        )
        db.commit()
        cursor.execute(
            "SELECT * FROM personas WHERE ci = %s",
            (persona.ci,),
        )
        return cursor.fetchone()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"CI {persona.ci} ya existe")
    finally:
        cursor.close()


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().lower().replace(" ", "_")


def _cell_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _parse_personas_workbook(content: bytes) -> tuple[list[PersonaCreate], list[dict[str, Any]]]:
    """
    Lee un .xlsx y devuelve filas válidas y errores por número de fila (Excel).
    """
    workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows, None)
    finally:
        workbook.close()

    if not header_row:
        raise HTTPException(status_code=400, detail="El Excel está vacío")

    headers = [_normalize_header(cell) for cell in header_row]
    required = {"nombre", "primer_apellido", "ci"}
    missing = required - set(headers)
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Faltan columnas obligatorias: {', '.join(sorted(missing))}",
        )

    idx = {name: headers.index(name) for name in headers if name}

    personas: list[PersonaCreate] = []
    errores: list[dict[str, Any]] = []
    excel_row = 1

    workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    try:
        for excel_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
                continue

            def col(name: str) -> Optional[str]:
                if name not in idx:
                    return None
                position = idx[name]
                if position >= len(row):
                    return None
                return _cell_text(row[position])

            payload = {
                "nombre": col("nombre"),
                "primer_apellido": col("primer_apellido"),
                "segundo_apellido": col("segundo_apellido"),
                "ci": col("ci"),
            }
            if not payload["nombre"] and not payload["primer_apellido"] and not payload["ci"]:
                continue

            try:
                personas.append(PersonaCreate(**payload))
            except ValidationError as exc:
                errores.append({"fila": excel_row, "detalle": exc.errors()})
    finally:
        workbook.close()

    return personas, errores


@app.post("/by-paul/personas/import-excel")
async def importar_personas_excel(
    archivo: UploadFile = File(..., description="Excel .xlsx con columnas de personas"),
    db=Depends(get_db),
):
    """
    Importa personas desde Excel.

    Encabezados esperados (fila 1): nombre, primer_apellido, segundo_apellido, ci.
    Las filas con CI duplicado se omiten; el resto se inserta en una transacción.
    """
    filename = (archivo.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Solo se aceptan archivos .xlsx")

    content = await archivo.read()
    if not content:
        raise HTTPException(status_code=400, detail="El archivo está vacío")

    personas, errores_validacion = _parse_personas_workbook(content)
    if not personas and errores_validacion:
        raise HTTPException(
            status_code=400,
            detail={"mensaje": "Ninguna fila válida", "errores": errores_validacion},
        )

    insertadas = 0
    omitidas = 0
    cursor = db.cursor(cursor_factory=RealDictCursor)
    try:
        for persona in personas:
            cursor.execute(
                """
                INSERT INTO personas (nombre, primer_apellido, segundo_apellido, ci)
                VALUES (%s, %s, %s, %s)
                ON CONFLICT (ci) DO NOTHING
                RETURNING persona_id
                """,
                (
                    persona.nombre,
                    persona.primer_apellido,
                    persona.segundo_apellido,
                    persona.ci,
                ),
            )
            if cursor.fetchone():
                insertadas += 1
            else:
                omitidas += 1
        db.commit()
    finally:
        cursor.close()

    return {
        "archivo": archivo.filename,
        "filas_validas": len(personas),
        "insertadas": insertadas,
        "omitidas_duplicado": omitidas,
        "errores_validacion": errores_validacion,
    }


@app.get("/by-paul/personas")
def listar_personas(db=Depends(get_db)):
    cursor = db.cursor(cursor_factory=RealDictCursor)
    cursor.execute("SELECT * FROM personas ORDER BY persona_id")
    resultados = cursor.fetchall()
    cursor.close()
    return resultados


@app.delete("/by-paul/personas/{persona_id}")
def eliminar_persona(persona_id: int, db=Depends(get_db)):
    cursor = db.cursor(cursor_factory=RealDictCursor)

    cursor.execute("SELECT * FROM personas WHERE persona_id = %s", (persona_id,))
    persona = cursor.fetchone()

    if not persona:
        cursor.close()
        raise HTTPException(status_code=404, detail="Persona no encontrada")

    cursor.execute("DELETE FROM personas WHERE persona_id = %s", (persona_id,))
    db.commit()
    cursor.close()

    return {
        "mensaje": "Persona eliminada",
        "persona": persona,
        "nota": "Sus viajes también fueron eliminados por CASCADE",
    }


# ============================================================================
# CRUD VIAJES
# ============================================================================

@app.post("/by-paul/viajes", status_code=201)
def crear_viaje(viaje: ViajeCreate, db=Depends(get_db)):
    cursor = db.cursor(cursor_factory=RealDictCursor)

    cursor.execute("SELECT persona_id FROM personas WHERE persona_id = %s", (viaje.persona_id,))
    if not cursor.fetchone():
        cursor.close()
        raise HTTPException(status_code=404, detail="Persona no existe")

    cursor.execute("SELECT pais_id FROM paises WHERE pais_id = %s", (viaje.pais_id,))
    if not cursor.fetchone():
        cursor.close()
        raise HTTPException(status_code=404, detail="País no existe")

    cursor.execute(
        """
        INSERT INTO viajes (persona_id, pais_id, fecha_llegada)
        VALUES (%s, %s, %s)
        RETURNING *
        """,
        (viaje.persona_id, viaje.pais_id, viaje.fecha_llegada),
    )
    db.commit()
    nuevo_viaje = cursor.fetchone()
    cursor.close()
    return nuevo_viaje


@app.get("/by-paul/viajes")
def listar_viajes(db=Depends(get_db)):
    cursor = db.cursor(cursor_factory=RealDictCursor)
    cursor.execute("""
        SELECT
            v.viaje_id,
            v.fecha_llegada,
            p.nombre || ' ' || p.primer_apellido AS persona,
            pa.pais_nombre AS pais
        FROM viajes v
        JOIN personas p ON v.persona_id = p.persona_id
        JOIN paises pa ON v.pais_id = pa.pais_id
        ORDER BY v.viaje_id
    """)
    resultados = cursor.fetchall()
    cursor.close()
    return resultados
